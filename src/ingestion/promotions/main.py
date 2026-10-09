from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import text
from src.simulators.pos_seed.connection import get_engine


PROJECT_ROOT = Path(__file__).resolve().parents[3]
PROMOTIONS_DIR = PROJECT_ROOT / "data" / "raw" / "promotions"

REQUIRED_COLUMNS = [
    "sku",
    "tienda_o_canal",
    "descuento_pct",
    "fecha_inicio",
    "fecha_fin",
]


def read_promotions(file_path: Path) -> list[dict]:
    """Read an Excel file and validate its promotion records."""
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if headers is None:
            raise ValueError(f"File {file_path.name} is empty.")

        headers = [
            str(value).strip() if value is not None else ""
            for value in headers
        ]

        if headers != REQUIRED_COLUMNS:
            raise ValueError(
                f"Invalid columns in {file_path.name}. "
                f"Expected: {REQUIRED_COLUMNS}. Received: {headers}"
            )

        promotions = []

        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue

            if len(row) != len(REQUIRED_COLUMNS):
                raise ValueError(
                    f"Incorrect number of columns in row {row_number}."
                )

            promotion = dict(zip(REQUIRED_COLUMNS, row))

            if any(value is None for value in promotion.values()):
                raise ValueError(
                    f"Empty fields detected in row {row_number}."
                )

            if not isinstance(promotion["sku"], str):
                raise ValueError(f"Invalid SKU in row {row_number}.")

            if not isinstance(promotion["tienda_o_canal"], str):
                raise ValueError(
                    f"Invalid store or channel in row {row_number}."
                )

            discount = promotion["descuento_pct"]

            if (
                not isinstance(discount, (int, float))
                or not 0 < discount <= 100
            ):
                raise ValueError(
                    f"Invalid discount in row {row_number}: {discount}"
                )

            start_date = promotion["fecha_inicio"]
            end_date = promotion["fecha_fin"]

            if not hasattr(start_date, "date") or not hasattr(end_date, "date"):
                raise ValueError(
                    f"Invalid dates in row {row_number}."
                )

            if start_date.date() > end_date.date():
                raise ValueError(
                    f"Start date is after end date in row {row_number}."
                )

            promotions.append(promotion)

        return promotions

    finally:
        workbook.close()

def insert_promotions(promotions: list[dict]) -> int:
    """Insert new promotions and skip existing records."""
    
    insert_query = text("""
        INSERT INTO public.promotions (
            sku,
            store_or_channel,
            discount_pct,
            start_date,
            end_date
        )
        SELECT
            CAST(:sku AS VARCHAR(50)),
            CAST(:store_or_channel AS VARCHAR(100)),
            CAST(:discount_pct AS NUMERIC(5,2)),
            CAST(:start_date AS DATE),
            CAST(:end_date AS DATE)
        WHERE NOT EXISTS (
            SELECT 1
            FROM public.promotions
            WHERE sku = CAST(:sku AS VARCHAR(50))
              AND store_or_channel = CAST(:store_or_channel AS VARCHAR(100))
              AND discount_pct = CAST(:discount_pct AS NUMERIC(5,2))
              AND start_date = CAST(:start_date AS DATE)
              AND end_date = CAST(:end_date AS DATE)
        )
    """)
    
    inserted_count = 0
    engine = get_engine()

    try:
        with engine.begin() as connection:
            for promotion in promotions:
                result = connection.execute(
                    insert_query,
                    {
                        "sku": promotion["sku"],
                        "store_or_channel": promotion["tienda_o_canal"],
                        "discount_pct": promotion["descuento_pct"],
                        "start_date": promotion["fecha_inicio"].date(),
                        "end_date": promotion["fecha_fin"].date(),
                    },
                )

                inserted_count += result.rowcount

        return inserted_count

    finally:
        engine.dispose()


def main():
    files = sorted(PROMOTIONS_DIR.glob("*.xlsx"))

    if not files:
        raise FileNotFoundError(
            f"No Excel files found in {PROMOTIONS_DIR}"
        )

    total_validated = 0
    total_inserted = 0

    for file_path in files:
        promotions = read_promotions(file_path)
        inserted = insert_promotions(promotions)

        total_validated += len(promotions)
        total_inserted += inserted

        print(
            f"{file_path.name}: "
            f"{len(promotions)} valid promotions, "
            f"{inserted} inserted, "
            f"{len(promotions) - inserted} skipped"
        )

    print(f"Total validated: {total_validated}")
    print(f"Total inserted: {total_inserted}")
    print(f"Total skipped: {total_validated - total_inserted}")


if __name__ == "__main__":
    main()
