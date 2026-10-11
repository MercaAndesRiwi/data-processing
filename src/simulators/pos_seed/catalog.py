from openpyxl import load_workbook
from sqlalchemy import text

from src.simulators.pos_seed.connection import get_engine
from src.simulators.pos_seed.config import COUNTRIES, PROJECT_ROOT


CATALOG_FILE = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "catalog"
    / "MercaAndes_Catalogo_Marketplace_Completo.xlsx"
)


def load_catalog():
    """Read active products and their base USD prices from the master catalog."""
    if not CATALOG_FILE.exists():
        raise FileNotFoundError(
            f"Master catalog not found: {CATALOG_FILE}"
        )

    workbook = load_workbook(
        CATALOG_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows)
        indexes = {name: i for i, name in enumerate(headers)}

        required_columns = {
            "SKU_ID",
            "Nombre_Producto",
            "Categoría",
            "País",
            "Estado",
            "Precio_Base_USD",
        }

        missing = required_columns - indexes.keys()
        if missing:
            raise ValueError(
                f"Missing catalog columns: {sorted(missing)}"
            )

        products_by_country = {code: [] for code in COUNTRIES}

        for row in rows:
            if not row or not row[indexes["SKU_ID"]]:
                continue

            country = str(row[indexes["País"]]).strip().upper()
            status = str(row[indexes["Estado"]]).strip().lower()

            # The POS simulator only operates in the five active countries.
            if country not in COUNTRIES:
                continue

            if status not in {"activo", "active"}:
                continue

            sku = str(row[indexes["SKU_ID"]]).strip()
            name = str(row[indexes["Nombre_Producto"]]).strip()
            category = str(row[indexes["Categoría"]]).strip()
            raw_price = row[indexes["Precio_Base_USD"]]

            if not name or not category:
                raise ValueError(
                    f"Product {sku} has no name or category."
                )

            if raw_price is None:
                raise ValueError(
                    f"Product {sku} has no base USD price."
                )

            base_price_usd = float(raw_price)

            if base_price_usd <= 0:
                raise ValueError(
                    f"Product {sku} has an invalid base USD price."
                )

            products_by_country[country].append(
                {
                    "sku": sku,
                    "name": name,
                    "category": category,
                    "base_price_usd": base_price_usd,
                }
            )

        return products_by_country

    finally:
        workbook.close()


def create_catalogs():
    products_by_country = load_catalog()
    engine = get_engine()

    try:
        with engine.begin() as connection:
            for country_code, schema_name in COUNTRIES.items():

                # Update existing databases as well as newly created ones.
                connection.execute(
                    text(
                        f"""
                        ALTER TABLE {schema_name}.product
                        ADD COLUMN IF NOT EXISTS
                        base_price_usd NUMERIC(12, 2)
                        """
                    )
                )

                products = products_by_country[country_code]
                print(
                    f"{country_code}: loading {len(products)} active products."
                )

                categories = sorted(
                    {product["category"] for product in products}
                )

                existing_categories = dict(
                    connection.execute(
                        text(
                            f"""
                            SELECT name, product_category_id
                            FROM {schema_name}.product_category
                            """
                        )
                    ).fetchall()
                )

                next_category_id = (
                    max(existing_categories.values(), default=0) + 1
                )

                for category_name in categories:
                    if category_name in existing_categories:
                        continue

                    connection.execute(
                        text(
                            f"""
                            INSERT INTO {schema_name}.product_category
                                (product_category_id, name)
                            VALUES (:category_id, :name)
                            """
                        ),
                        {
                            "category_id": next_category_id,
                            "name": category_name,
                        },
                    )

                    existing_categories[category_name] = next_category_id
                    next_category_id += 1

                for product in products:
                    connection.execute(
                        text(
                            f"""
                            INSERT INTO {schema_name}.product
                                (
                                    product_id,
                                    product_category_id,
                                    sku,
                                    name,
                                    base_price_usd
                                )
                            VALUES
                                (
                                    :product_id,
                                    :category_id,
                                    :sku,
                                    :name,
                                    :base_price_usd
                                )
                            ON CONFLICT (sku) DO UPDATE SET
                                product_category_id = EXCLUDED.product_category_id,
                                name = EXCLUDED.name,
                                base_price_usd = EXCLUDED.base_price_usd
                            """
                        ),
                        {
                            "product_id": product["sku"],
                            "category_id": existing_categories[
                                product["category"]
                            ],
                            "sku": product["sku"],
                            "name": product["name"],
                            "base_price_usd": product["base_price_usd"],
                        },
                    )

                print(
                    f"{country_code}: catalog loaded successfully."
                )

        print("Master catalog loading completed.")

    finally:
        engine.dispose()


if __name__ == "__main__":
    create_catalogs()
